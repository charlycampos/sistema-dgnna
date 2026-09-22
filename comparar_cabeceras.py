import openpyxl

wb_old = openpyxl.load_workbook('docs/6. CONSOLIDADO_RENE_2018-2025_A JUNIO2026.xlsx', read_only=True)
ws_old = wb_old['Hoja1']
header_old = None
for i, r in enumerate(ws_old.iter_rows(min_row=6, max_row=6, values_only=True)):
    header_old = r
    break

wb_new = openpyxl.load_workbook('docs/RENE PROTECCIÓN ESPECIAL 2026.xlsx', read_only=True)
ws_new = wb_new['RENE MATRIZ']
header_new = None
for i, r in enumerate(ws_new.iter_rows(min_row=6, max_row=6, values_only=True)):
    header_new = r
    break

print(f"Header Old len: {len(header_old)}, Header New len: {len(header_new)}")
diffs = []
for idx in range(min(len(header_old), len(header_new))):
    old_val = str(header_old[idx]).strip() if header_old[idx] is not None else ""
    new_val = str(header_new[idx]).strip() if header_new[idx] is not None else ""
    if old_val != new_val:
        diffs.append((idx, old_val, new_val))

print(f"Diferencias encontradas: {len(diffs)}")
for d in diffs[:20]:
    print(f"Col {d[0]}: OLD='{d[1]}' vs NEW='{d[2]}'")
