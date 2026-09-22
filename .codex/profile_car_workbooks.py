from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
FILES = [
    p for p in (ROOT / "docs").glob("*.xlsx")
    if "RENE" in p.name.upper() and ("CAR" in p.name.upper() or "NNA" in p.name.upper())
]


def clean(value):
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).strip())
    return text or None


def normalized(value):
    text = clean(value) or ""
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", "_", text.upper()).strip("_")


def category_safe(header):
    h = normalized(header)
    blocked = ("NOMBRE", "APELLIDO", "DNI", "DOCUMENTO", "DIRECCION", "DOMICILIO", "TELEFONO", "CORREO")
    return not any(token in h for token in blocked)


def infer_type(values):
    vals = [v for v in values if v not in (None, "")]
    if not vals:
        return "vacio"
    kinds = Counter()
    for value in vals:
        if isinstance(value, bool): kinds["booleano"] += 1
        elif hasattr(value, "year") and hasattr(value, "month"): kinds["fecha"] += 1
        elif isinstance(value, int): kinds["entero"] += 1
        elif isinstance(value, float): kinds["decimal"] += 1
        else: kinds["texto"] += 1
    kind, count = kinds.most_common(1)[0]
    return kind if count == len(vals) else "mixto:" + ",".join(f"{k}={v}" for k, v in kinds.items())


def profile_sheet(ws):
    preview = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 40), values_only=True))
    if not preview:
        return {
            "hoja": ws.title,
            "fila_encabezado": None,
            "filas_datos": 0,
            "columnas": 0,
            "encabezados_duplicados": [],
            "variables": [],
        }
    scores = [sum(clean(v) is not None for v in row) for row in preview]
    header_idx = max(range(len(scores)), key=lambda i: scores[i]) if scores else 0
    headers = [clean(v) or f"SIN_ENCABEZADO_{i+1}" for i, v in enumerate(preview[header_idx])]
    # Trim empty trailing worksheet columns.
    last = max((i for i, h in enumerate(headers) if not h.startswith("SIN_ENCABEZADO_")), default=-1)
    headers = headers[: last + 1]
    data = []
    for row in ws.iter_rows(min_row=header_idx + 2, values_only=True):
        vals = list(row[:len(headers)])
        if any(clean(v) is not None for v in vals):
            data.append(vals)
    cols = []
    for idx, header in enumerate(headers):
        values = [row[idx] if idx < len(row) else None for row in data]
        non_null = [v for v in values if clean(v) is not None]
        strings = [clean(v) for v in non_null]
        counts = Counter(strings)
        item = {
            "posicion": idx + 1,
            "encabezado": header,
            "nombre_normalizado": normalized(header),
            "tipo_inferido": infer_type(non_null),
            "no_nulos": len(non_null),
            "nulos": len(data) - len(non_null),
            "unicos": len(counts),
            "duplicados_no_nulos": sum(v - 1 for v in counts.values() if v > 1),
        }
        if category_safe(header) and 0 < len(counts) <= 25:
            item["categorias"] = [{"valor": k, "conteo": v} for k, v in counts.most_common(25)]
        cols.append(item)
    duplicate_headers = [h for h, n in Counter(normalized(h) for h in headers).items() if n > 1]
    return {
        "hoja": ws.title,
        "fila_encabezado": header_idx + 1,
        "filas_datos": len(data),
        "columnas": len(headers),
        "encabezados_duplicados": duplicate_headers,
        "variables": cols,
    }


result = []
for path in sorted(FILES):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    result.append({
        "archivo": path.name,
        "tamano_bytes": path.stat().st_size,
        "hojas": [profile_sheet(ws) for ws in wb.worksheets],
    })
    wb.close()

output = ROOT / "profile_car_workbooks.json"
output.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
print(json.dumps({
    "reporte": str(output),
    "archivos": [
        {
            "archivo": book["archivo"],
            "hojas": [
                {"hoja": sheet["hoja"], "filas": sheet["filas_datos"], "columnas": sheet["columnas"]}
                for sheet in book["hojas"]
            ],
        }
        for book in result
    ],
}, ensure_ascii=False, indent=2))
