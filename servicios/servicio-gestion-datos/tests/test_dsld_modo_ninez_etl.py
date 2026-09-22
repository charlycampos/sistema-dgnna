"""
Pruebas del ETL DSLD · Ponte en Modo Niñez.

- Pruebas con datos sintéticos (siempre se ejecutan).
- Prueba de referencia contra docs/MATRIZ DE REPORTE PBI 2026.xlsx (se omite si no está).
  Cifras esperadas = medidas MEDIDAS_MODO_NIÑEZ del Power BI con la matriz del 28/02/2026.

Ejecutar:  cd servicios/servicio-gestion-datos && python -m pytest tests -q
"""

import io
import os
import sys
from datetime import datetime

import openpyxl
import pandas as pd
import pytest
from openpyxl.worksheet.table import Table

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from domain.services.dsld_excel_reader import ExcelLecturaError, leer_tabla_excel  # noqa: E402
from domain.services.dsld_modo_ninez_etl import (  # noqa: E402
    COLUMNAS_EXCLUIDAS,
    COLUMNAS_REQUERIDAS,
    COLUMNAS_USADAS,
    TABLA_EXCEL,
    EtlModoNinezError,
    transformar_modo_ninez,
)

RUTA_XLSX = os.getenv(
    "DSLD_MODO_NINEZ_PRUEBA",
    os.path.join(os.path.dirname(__file__), "..", "..", "..", "docs", "MATRIZ DE REPORTE PBI 2026.xlsx"),
)

ENCABEZADOS = ["Nº", "UBIGEO", "NOMBRE DE GOBIERNO", "TIPO DE GOBIERNO", "MACROREGIÓN", "DEPARTAMENTO",
               "PROVINCIA", "DISTRITO", "MODO_NIÑEZ", "AÑO QUE SE SUMÓ A LA ESTRATEGIA",
               "FECHA DE PRESENTACIÓN", "FECHA DE ACTA DE COMPROMISO", "SR/SRA", "ALCALDE/SA",
               "CÓDIGO DE DEMUNA", "ESTADO"]


def _filas():
    return [
        [1, "150000", "Gobierno Regional de Lima", "Regional", "Lima y Callao", "Lima", None, None, "SI", 2021,
         datetime(2026, 2, 10), None, "Sr.", "Persona Ficticia Uno", None, None],
        [2, "150101", "Municipalidad Metropolitana de Lima", "Provincial", "Lima y Callao", "Lima", "Lima", "Lima",
         "SI", 2019, None, "29/02/2023", "Sra.", "Persona Ficticia Dos", 15001, "Acreditada"],
        [3, 10201, "Municipalidad Provincial de Bagua", "Provincial", "Norte", "Amazonas", "Bagua", "Bagua",
         "SI", 2023, None, datetime(2023, 5, 2), "Sr.", "Persona Ficticia Tres", 1004, "NO ACREDITADA"],
        [4, "050101", "Municipalidad Distrital de Ayacucho", "Distrital", "Centro oriente", "Ayacucho",
         "Huamanga", "Ayacucho", "NO", None, None, None, "Sr.", "Persona Ficticia Cuatro", "05001", "Acreditada"],
        [None, None, None, None, None, None, None, None, None, None, None, None, None, None, None, None],
    ]


def _df(filas=None):
    return pd.DataFrame(filas if filas is not None else _filas(), columns=ENCABEZADOS, dtype=object)


def _excel(filas=None, nombre_tabla=TABLA_EXCEL) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nota de la hoja fuera de la tabla"])
    ws.append(ENCABEZADOS)
    datos = filas if filas is not None else _filas()[:-1]
    for f in datos:
        ws.append(f)
    ultima_col = openpyxl.utils.get_column_letter(len(ENCABEZADOS))
    ws.add_table(Table(displayName=nombre_tabla, ref=f"A2:{ultima_col}{len(datos) + 2}"))
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def test_reglas_basicas():
    r = transformar_modo_ninez(_df())
    g = {x["ubigeo"]: x for x in r.gobiernos}
    assert len(r.gobiernos) == 4 and r.descartados == 1
    assert g["150000"]["nivel_gobierno"] == "REGIONAL" and g["150000"]["departamento_mod"] == "GORE LIMA"
    assert g["150000"]["ubigeo_prov"] is None and g["150000"]["anio_presentacion"] == 2026
    assert g["150101"]["departamento_mod"] == "LIMA METROPOLITANA"
    assert g["010201"]["ubigeo_prov"] == "0102"                 # ubigeo numérico con cero perdido
    assert g["010201"]["codigo_demuna"] == "01004"
    assert g["010201"]["estado_demuna"] == "NO ACREDITADA"
    assert g["150101"]["fecha_acta"] is None                     # 29/02/2023 no existe
    assert any("29/02/2023" in a for a in r.advertencias)
    assert g["050101"]["adherido"] == "N"
    res = r.resumen()
    assert (res["adheridos"], res["regionales"], res["provinciales"], res["distritales"]) == (3, 1, 2, 0)
    assert res["presentaronReporte"] == 1


def test_no_guarda_datos_personales():
    r = transformar_modo_ninez(_df())
    claves = set().union(*(x.keys() for x in r.gobiernos))
    assert not claves & {"sr_sra", "alcalde", "alcalde_sa", "SR/SRA", "ALCALDE/SA"}
    valores = {str(v) for x in r.gobiernos for v in x.values()}
    assert not any("Persona Ficticia" in v for v in valores)


def test_lector_excel_solo_columnas_usadas():
    df = leer_tabla_excel(_excel(), TABLA_EXCEL, COLUMNAS_USADAS, COLUMNAS_REQUERIDAS)
    assert len(df) == 4
    assert not set(COLUMNAS_EXCLUIDAS) & set(df.columns)
    assert transformar_modo_ninez(df).resumen()["adheridos"] == 3


def test_lector_excel_sin_tabla():
    with pytest.raises(ExcelLecturaError, match=TABLA_EXCEL):
        leer_tabla_excel(_excel(nombre_tabla="OTRA_TABLA"), TABLA_EXCEL)
    with pytest.raises(ExcelLecturaError, match="Excel"):
        leer_tabla_excel(b"no es un excel", TABLA_EXCEL)


def test_errores_no_cargan_nada():
    filas = _filas()
    filas[1][3] = "Comunal"
    filas[2][8] = "QUIZÁS"
    filas[3][1] = "150000"          # repetido
    with pytest.raises(EtlModoNinezError) as e:
        transformar_modo_ninez(_df(filas))
    msg = str(e.value)
    assert "3 fila(s)" in msg and "Comunal" in msg and "QUIZÁS" in msg and "repetido" in msg


def test_falta_columna():
    with pytest.raises(EtlModoNinezError, match="MODO_NIÑEZ"):
        transformar_modo_ninez(_df().drop(columns=["MODO_NIÑEZ"]))


def test_catalogo_ubigeo():
    cat = {"010200": {"departamento": "AMAZONAS", "provincia": "BAGUA", "distrito": None},
           "150000": {"departamento": "LIMA", "provincia": None, "distrito": None}}
    g = {x["ubigeo"]: x for x in transformar_modo_ninez(_df(), cat).gobiernos}
    assert g["010201"]["provincia"] == "BAGUA" and g["010201"]["distrito"] is None
    assert g["150000"]["departamento"] == "LIMA"


@pytest.mark.skipif(not os.path.isfile(RUTA_XLSX), reason="MATRIZ DE REPORTE PBI 2026.xlsx no disponible")
def test_referencia_power_bi():
    df = leer_tabla_excel(RUTA_XLSX, TABLA_EXCEL, COLUMNAS_USADAS, COLUMNAS_REQUERIDAS)
    r = transformar_modo_ninez(df).resumen()
    assert (r["gobiernos"], r["adheridos"]) == (546, 545)
    assert (r["regionales"], r["provinciales"], r["distritales"]) == (15, 136, 394)
    assert r["presentaronReporte"] == 71
    assert r["presentaronPorNivel"] == {"REGIONAL": 4, "PROVINCIAL": 10, "DISTRITAL": 57}
    assert r["adheridosPorAnio"] == {2019: 51, 2020: 33, 2021: 62, 2022: 42, 2023: 121, 2024: 105, 2025: 71, 2026: 45}
