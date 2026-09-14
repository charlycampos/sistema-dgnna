"""Pruebas unitarias para el motor de renderizado DOCX institucional."""

import io
import json
import os
import sys
import unittest
from datetime import datetime
from pathlib import Path
from docx import Document

SERVICE_DIR = Path(__file__).resolve().parents[1]
os.environ["TESTING"] = "true"
os.environ["DATABASE_URL"] = "sqlite:///:memory:"
sys.path.insert(0, str(SERVICE_DIR))

from domain.services.docx_renderer import renderizar  # noqa: E402


class DummyObject:
    def __init__(self, **kwargs):
        for k, v in kwargs.items():
            setattr(self, k, v)


class DocxRendererTest(unittest.TestCase):
    def setUp(self):
        self.plantilla = DummyObject(
            id="plt-1",
            nombre="Ayuda Memoria de Servicios NNA",
            secciones=[
                DummyObject(
                    id="sec-1",
                    orden=1,
                    titulo="Resumen Ejecutivo",
                    tipoSeccion="TEXTO"
                ),
                DummyObject(
                    id="sec-2",
                    orden=2,
                    titulo="Cobertura de Servicios",
                    tipoSeccion="TABLA_DATOS"
                ),
                DummyObject(
                    id="sec-3",
                    orden=3,
                    titulo="Indicador Clave",
                    tipoSeccion="KPI"
                )
            ]
        )

        self.valores = [
            DummyObject(
                seccionId="sec-1",
                textoContenido="Este es un resumen ejecutivo oficial del Sistema DGNNA.\n- Primer punto de atención desplegado.\n- Segundo punto relevante.",
                datosTablaJson=None,
                cifraCorte=None
            ),
            DummyObject(
                seccionId="sec-2",
                textoContenido="",
                datosTablaJson=json.dumps([
                    {"Departamento": "Lima", "Atenciones": "1,250", "Cobertura": "88%"},
                    {"Departamento": "Cusco", "Atenciones": "450", "Cobertura": "72%"}
                ]),
                cifraCorte=None
            ),
            DummyObject(
                seccionId="sec-3",
                textoContenido="",
                datosTablaJson=None,
                cifraCorte="1,700 NNA Atendidos"
            )
        ]

        self.acciones = [
            DummyObject(
                fecha="11/09/2026",
                institucion="UPE Lima",
                descripcion="Medida de protección provisional dictada con éxito."
            )
        ]

    def test_renderizar_documento_borrador(self):
        doc_obj = DummyObject(
            codigoInterno="AM-2026-001",
            titulo="Informe Situacional de Servicios NNA",
            direccion="DPE",
            region="Lima",
            fechaCorte="Septiembre 2026",
            estado="BORRADOR",
            nivelRiesgo="ALTO",
            servicioMimp="UPE",
            versionDoc="1.0",
            documentoOrigenId=None,
            creadoPor="Especialista DPE",
            publicadoPor=None,
            publicadoAt=None,
            hashIntegridad=None,
            plantilla=self.plantilla,
            valores=self.valores,
            acciones=self.acciones
        )

        stream = renderizar(doc_obj)
        self.assertIsInstance(stream, io.BytesIO)
        self.assertGreater(stream.getbuffer().nbytes, 0)

        # Verificar que python-docx puede reabrir el documento generado
        docx_doc = Document(stream)
        textos = [p.text for p in docx_doc.paragraphs]
        contenido_completo = " ".join(textos)

        self.assertIn("AYUDA MEMORIA INSTITUCIONAL", contenido_completo)
        self.assertIn("INFORME SITUACIONAL DE SERVICIOS NNA", contenido_completo)
        self.assertIn("1. RESUMEN EJECUTIVO", contenido_completo)
        self.assertIn("Este es un resumen ejecutivo oficial", contenido_completo)
        self.assertIn("CRONOLOGÍA DE ACCIONES", contenido_completo)

        # Tablas presentes (Metadatos, Datos, Bitácora)
        self.assertGreaterEqual(len(docx_doc.tables), 3)

    def test_renderizar_documento_publicado_con_hash(self):
        doc_pub = DummyObject(
            codigoInterno="AM-2026-002",
            titulo="Reporte Final de Medidas",
            direccion="DSNNA",
            region="Nacional",
            fechaCorte="Septiembre 2026",
            estado="PUBLICADO",
            nivelRiesgo="MODERADO",
            servicioMimp="DEMUNA",
            versionDoc="1.0",
            documentoOrigenId=None,
            creadoPor="Especialista DSNNA",
            publicadoPor="Dirección DSNNA",
            publicadoAt=datetime(2026, 9, 11, 12, 0),
            hashIntegridad="e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
            plantilla=self.plantilla,
            valores=self.valores,
            acciones=self.acciones
        )

        stream = renderizar(doc_pub)
        docx_doc = Document(stream)
        
        # Debe contener la tabla de sello criptográfico SHA-256
        tabla_textos = []
        for t in docx_doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    tabla_textos.append(cell.text)
        
        texto_tablas = " ".join(tabla_textos)
        self.assertIn("DOCUMENTO OFICIAL PUBLICADO E INMUTABLE", texto_tablas)
        self.assertIn("e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855", texto_tablas)


if __name__ == "__main__":
    unittest.main()
